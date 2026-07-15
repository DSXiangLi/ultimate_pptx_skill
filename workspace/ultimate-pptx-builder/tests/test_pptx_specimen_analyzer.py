#!/usr/bin/env python3
"""Tests for PPTX specimen analyzer / clone C1 evidence pack."""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT as PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ANALYZER = ROOT / "scripts" / "analyze_pptx_specimen.py"
sys.path.insert(0, str(ROOT / "scripts"))


def make_sample_png(path: Path) -> None:
    # 1x1 transparent PNG
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
        b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\x0bIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def make_sample_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(5.5), Inches(0.8))
    title.name = "sample_title"
    title.text_frame.text = "Clone Analyzer Title"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.6), Inches(3.0), Inches(1.2))
    panel.name = "sample_panel"
    panel.text = "Metric 42%"

    img = path.with_suffix(".png")
    make_sample_png(img)
    pic = slide.shapes.add_picture(str(img), Inches(5.8), Inches(1.5), Inches(1.0), Inches(1.0))
    pic.name = "sample_picture"

    group = slide.shapes.add_group_shape()
    group.name = "sample_group"
    grouped = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.4), Inches(2.2), Inches(0.55))
    grouped.name = "sample_grouped_metric"
    grouped.text = "Grouped Insight"

    empty = slide.shapes.add_textbox(Inches(4.2), Inches(3.4), Inches(1.2), Inches(0.4))
    empty.name = "sample_empty_text_container"

    table_shape = slide.shapes.add_table(2, 2, Inches(4.2), Inches(4.1), Inches(2.4), Inches(0.9))
    table_shape.name = "sample_native_table"
    table_shape.table.cell(0, 0).text = "Driver"
    table_shape.table.cell(0, 1).text = "Value"
    table_shape.table.cell(1, 0).text = "Speed"
    table_shape.table.cell(1, 1).text = "High"
    for c in range(2):
        cell = table_shape.table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x1A, 0x2B, 0x3C)
        cell.text_frame.margin_left = 91440
        cell.text_frame.margin_right = 91440
        cell.text_frame.margin_top = 45720
        cell.text_frame.margin_bottom = 45720
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.runs[0]
        run.font.name = "Aptos Display"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for c in range(2):
        cell = table_shape.table.cell(1, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xDD, 0xEE, 0xFF)
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.runs[0]
        run.font.name = "Aptos"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)

    prs.save(str(path))


class PptxSpecimenAnalyzerTests(unittest.TestCase):
    def test_contact_sheet_is_created_from_rendered_slide_images(self):
        from analyze_pptx_specimen import create_contact_sheet

        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            rendered = tmp / "rendered"
            rendered.mkdir()
            Image.new("RGB", (120, 80), (20, 30, 40)).save(rendered / "slide-01.png")
            Image.new("RGB", (120, 80), (60, 70, 80)).save(rendered / "slide-02.png")

            sheet = create_contact_sheet(rendered)

            self.assertEqual(sheet, rendered / "contact-sheet.png")
            self.assertTrue(sheet.exists())
            with Image.open(sheet) as img:
                self.assertGreaterEqual(img.width, 240)
                self.assertGreaterEqual(img.height, 80)

    def test_analyzer_creates_evidence_pack_with_inventory_and_provenance(self):
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            pptx = tmp / "sample.pptx"
            out = tmp / "specimen"
            make_sample_pptx(pptx)

            proc = subprocess.run(
                [
                    sys.executable,
                    str(ANALYZER),
                    str(pptx),
                    "--out",
                    str(out),
                    "--deck-id",
                    "sample-deck",
                    "--source-url",
                    "https://example.test/template",
                    "--license-note",
                    "unit-test license note",
                    "--skip-render",
                ],
                cwd=str(ROOT),
                text=True,
                capture_output=True,
            )
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)

            expected_files = [
                "original.pptx",
                "provenance.json",
                "extracted-text.md",
                "object-inventory.json",
                "theme-inventory.json",
                "master-layout-inventory.json",
                "asset-inventory.json",
            ]
            for name in expected_files:
                self.assertTrue((out / name).exists(), f"missing {name}")
            self.assertTrue((out / "unpacked" / "ppt" / "slides" / "slide1.xml").exists())
            self.assertFalse((out / "rendered").exists(), "--skip-render should not create rendered output")

            provenance = json.loads((out / "provenance.json").read_text(encoding="utf-8"))
            self.assertEqual(provenance["deck_id"], "sample-deck")
            self.assertEqual(provenance["source_url"], "https://example.test/template")
            self.assertEqual(provenance["license_note"], "unit-test license note")
            self.assertEqual(provenance["slide_count"], 1)
            self.assertRegex(provenance["sha256"], r"^[0-9a-f]{64}$")

            text = (out / "extracted-text.md").read_text(encoding="utf-8")
            self.assertIn("Clone Analyzer Title", text)
            self.assertIn("Metric 42%", text)

            inventory = json.loads((out / "object-inventory.json").read_text(encoding="utf-8"))
            self.assertEqual(inventory["slide_count"], 1)
            self.assertGreaterEqual(inventory["object_count"], 3)
            texts = "\n".join(obj.get("text", "") for slide in inventory["slides"] for obj in slide["objects"])
            self.assertIn("Clone Analyzer Title", texts)
            self.assertIn("Metric 42%", texts)
            object_types = {obj["type"] for slide in inventory["slides"] for obj in slide["objects"]}
            self.assertIn("text", object_types)
            self.assertIn("shape", object_types)
            self.assertIn("image", object_types)
            first_obj = inventory["slides"][0]["objects"][0]
            self.assertIn("box", first_obj)
            self.assertIn("source_xml_path", first_obj)
            self.assertIn("source_shape_id", first_obj)
            self.assertIn("z", first_obj)

            theme = json.loads((out / "theme-inventory.json").read_text(encoding="utf-8"))
            self.assertGreaterEqual(theme["theme_count"], 1)
            self.assertTrue(theme["themes"][0]["colors"])

            masters = json.loads((out / "master-layout-inventory.json").read_text(encoding="utf-8"))
            self.assertGreaterEqual(masters["master_count"], 1)
            self.assertGreaterEqual(masters["layout_count"], 1)

            assets = json.loads((out / "asset-inventory.json").read_text(encoding="utf-8"))
            self.assertGreaterEqual(assets["media_count"], 1)
            self.assertTrue(any(item["path"].startswith("ppt/media/") for item in assets["media"]))

            with zipfile.ZipFile(out / "original.pptx") as zf:
                self.assertIn("ppt/slides/slide1.xml", zf.namelist())


if __name__ == "__main__":
    unittest.main()
