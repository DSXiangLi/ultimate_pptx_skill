#!/usr/bin/env python3
"""Analyze an existing PPTX specimen into a clone-learning evidence pack.

Phase C1 of the PPTX cloner reverse compiler:
- preserve provenance and original package
- unpack raw OOXML
- render slides when available
- extract text
- inventory objects, theme, masters/layouts, and assets
"""
from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import json
import posixpath
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
try:
    from PIL import Image, ImageDraw
except Exception:  # pragma: no cover - Pillow is expected in this project, but keep CLI robust.
    Image = None
    ImageDraw = None

ROOT = Path(__file__).resolve().parents[1]
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, data: Dict[str, Any]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def clean_dir(path: Path) -> None:
    if path.exists():
        shutil.rmtree(str(path))
    path.mkdir(parents=True, exist_ok=True)


def unzip_pptx(pptx: Path, outdir: Path) -> List[str]:
    clean_dir(outdir)
    with zipfile.ZipFile(str(pptx)) as zf:
        zf.extractall(str(outdir))
        return sorted(zf.namelist())


def emu_to_px(value: Any, slide_emu: int, base_px: int) -> float:
    if not slide_emu:
        return 0.0
    return round(float(value) / float(slide_emu) * float(base_px), 3)


def shape_type(shape: Any) -> str:
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if st == MSO_SHAPE_TYPE.CHART:
        return "chart"
    if st == MSO_SHAPE_TYPE.TABLE:
        return "table"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(shape, "has_text_frame", False) and st == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text"
    return "shape"


def resolve_package_target(source_part: str, target: str) -> str:
    """Resolve an OOXML relationship target to a package-relative path."""
    target = (target or "").replace("\\", "/")
    if not target:
        return ""
    if target.startswith("/"):
        return target.lstrip("/")
    base = posixpath.dirname(source_part.replace("\\", "/"))
    return posixpath.normpath(posixpath.join(base, target))


def picture_relationships(pptx: Path, slide_index: int) -> Dict[str, Dict[str, str]]:
    """Map picture cNvPr IDs on a slide to their image relationship target."""
    slide_part = "ppt/slides/slide{}.xml".format(slide_index)
    rels_part = "ppt/slides/_rels/slide{}.xml.rels".format(slide_index)
    try:
        with zipfile.ZipFile(str(pptx)) as zf:
            names = set(zf.namelist())
            slide_xml = zf.read(slide_part)
            rels_xml = zf.read(rels_part) if rels_part in names else b""
    except Exception:
        return {}

    rels: Dict[str, Dict[str, str]] = {}
    if rels_xml:
        try:
            rel_root = ET.fromstring(rels_xml)
            for rel in rel_root:
                rid = rel.attrib.get("Id", "")
                target = rel.attrib.get("Target", "")
                if rid:
                    rels[rid] = {
                        "target": target,
                        "type": rel.attrib.get("Type", ""),
                        "package_path": resolve_package_target(slide_part, target),
                    }
        except Exception:
            rels = {}

    try:
        root = ET.fromstring(slide_xml)
    except Exception:
        return {}
    out: Dict[str, Dict[str, str]] = {}
    for pic in root.findall(".//{}pic".format(P_NS)):
        c_nv_pr = pic.find(".//{}cNvPr".format(P_NS))
        blip = pic.find(".//{}blip".format(A_NS))
        if c_nv_pr is None or blip is None:
            continue
        shape_id = c_nv_pr.attrib.get("id", "")
        rid = blip.attrib.get(R_NS + "embed") or blip.attrib.get(R_NS + "link") or ""
        if shape_id and rid:
            rel = rels.get(rid, {})
            out[shape_id] = {
                "relationship_id": rid,
                "relationship_type": rel.get("type", ""),
                "target": rel.get("target", ""),
                "package_path": rel.get("package_path", ""),
            }
    for sp in root.findall(".//{}sp".format(P_NS)):
        c_nv_pr = sp.find(".//{}cNvPr".format(P_NS))
        blip = sp.find(".//{}blip".format(A_NS))
        if c_nv_pr is None or blip is None:
            continue
        shape_id = c_nv_pr.attrib.get("id", "")
        rid = blip.attrib.get(R_NS + "embed") or blip.attrib.get(R_NS + "link") or ""
        if shape_id and rid:
            rel = rels.get(rid, {})
            out[shape_id] = {
                "relationship_id": rid,
                "relationship_type": rel.get("type", ""),
                "target": rel.get("target", ""),
                "package_path": rel.get("package_path", ""),
            }
    return out


def slide_background(pptx: Path, slide_index: int) -> Dict[str, Any]:
    """Extract slide-level background material from p:cSld/p:bg.

    C3.2 starts with solid RGB backgrounds because the observed visual defect is
    dark/colored slides collapsing to white during rebuild. Other background
    classes are classified for later compiler increments instead of being
    silently ignored.
    """
    slide_part = "ppt/slides/slide{}.xml".format(slide_index)
    try:
        with zipfile.ZipFile(str(pptx)) as zf:
            root = ET.fromstring(zf.read(slide_part))
    except Exception:
        return {"kind": "unknown", "source_xml_path": slide_part}

    bg = root.find("{}cSld/{}bg".format(P_NS, P_NS))
    if bg is None:
        return {"kind": "default", "source_xml_path": slide_part}
    bg_pr = bg.find("{}bgPr".format(P_NS))
    if bg_pr is None:
        bg_ref = bg.find("{}bgRef".format(P_NS))
        return {"kind": "reference" if bg_ref is not None else "unknown", "source_xml_path": slide_part}
    solid = bg_pr.find("{}solidFill".format(A_NS))
    if solid is not None:
        rgb_val = hex_color_from_node(solid)
        if rgb_val:
            return {"kind": "solid", "rgb": rgb_val.upper(), "source_xml_path": slide_part}
        return {"kind": "solid-unresolved", "source_xml_path": slide_part}
    if bg_pr.find("{}gradFill".format(A_NS)) is not None:
        return {"kind": "gradient", "source_xml_path": slide_part, "degradation": "gradient background not yet reconstructed"}
    if bg_pr.find("{}blipFill".format(A_NS)) is not None:
        return {"kind": "image", "source_xml_path": slide_part, "degradation": "image background not yet reconstructed"}
    return {"kind": "unknown", "source_xml_path": slide_part}


def shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        try:
            return shape.text or ""
        except Exception:
            return ""
    return ""


def style_summary(shape: Any) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    try:
        fill = shape.fill
        out["fill_type"] = str(fill.type)
        if getattr(fill, "fore_color", None) is not None and fill.fore_color.rgb:
            out["fill_rgb"] = str(fill.fore_color.rgb)
    except Exception:
        pass
    try:
        line = shape.line
        if getattr(line, "color", None) is not None and line.color.rgb:
            out["line_rgb"] = str(line.color.rgb)
        if line.width:
            out["line_width_emu"] = int(line.width)
    except Exception:
        pass
    return out


def font_summary(shape: Any) -> Dict[str, Any]:
    result: Dict[str, Any] = {}
    if not getattr(shape, "has_text_frame", False):
        return result
    try:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                font = run.font
                if font.name and "name" not in result:
                    result["name"] = font.name
                if font.size and "size_pt" not in result:
                    result["size_pt"] = round(float(font.size.pt), 3)
                if font.bold is not None and "bold" not in result:
                    result["bold"] = bool(font.bold)
                if getattr(font.color, "rgb", None) and "color_rgb" not in result:
                    result["color_rgb"] = str(font.color.rgb)
    except Exception:
        pass
    return result


def table_summary(shape: Any) -> Dict[str, Any]:
    if not getattr(shape, "has_table", False):
        return {}
    try:
        table = shape.table
        cells: List[List[str]] = []
        cell_styles: List[List[Dict[str, Any]]] = []
        for row in table.rows:
            cell_texts: List[str] = []
            row_styles: List[Dict[str, Any]] = []
            for cell in row.cells:
                cell_texts.append(cell.text)
                style: Dict[str, Any] = {}
                try:
                    if str(cell.fill.type).startswith("SOLID") and getattr(cell.fill.fore_color, "rgb", None):
                        style["fill_rgb"] = str(cell.fill.fore_color.rgb)
                except Exception:
                    pass
                try:
                    tf = cell.text_frame
                    for key in ["margin_left", "margin_right", "margin_top", "margin_bottom"]:
                        value = getattr(tf, key, None)
                        if value is not None:
                            style[key + "_emu"] = int(value)
                    if tf.paragraphs:
                        para = tf.paragraphs[0]
                        if para.alignment is not None:
                            style["alignment"] = str(para.alignment).split()[0]
                        for run in para.runs:
                            font: Dict[str, Any] = {}
                            if run.font.name:
                                font["name"] = run.font.name
                            if run.font.size:
                                font["size_pt"] = round(float(run.font.size.pt), 3)
                            if run.font.bold is not None:
                                font["bold"] = bool(run.font.bold)
                            if run.font.italic is not None:
                                font["italic"] = bool(run.font.italic)
                            try:
                                if getattr(run.font.color, "rgb", None):
                                    font["color_rgb"] = str(run.font.color.rgb)
                            except Exception:
                                pass
                            if font:
                                style["font"] = font
                                break
                except Exception:
                    pass
                row_styles.append(style)
            cells.append(cell_texts)
            cell_styles.append(row_styles)
        return {
            "row_count": len(table.rows),
            "column_count": len(table.columns),
            "cells": cells,
            "cell_styles": cell_styles,
        }
    except Exception:
        return {}


def inventory_objects(pptx: Path) -> Dict[str, Any]:
    prs = Presentation(str(pptx))
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    slides: List[Dict[str, Any]] = []
    object_count = 0

    def object_from_shape(slide_index: int, shape: Any, z: float, image_rels: Dict[str, Dict[str, str]], parent_group: Optional[Dict[str, Any]] = None, depth: int = 0) -> Dict[str, Any]:
        text = shape_text(shape)
        obj_id = "slide{:02d}_shape{:03d}".format(slide_index, int(shape.shape_id))
        obj = {
            "id": obj_id,
            "name": getattr(shape, "name", ""),
            "type": shape_type(shape),
            "shape_type": str(shape.shape_type),
            "z": z,
            "box": {
                "x_px": emu_to_px(shape.left, slide_w, 1280),
                "y_px": emu_to_px(shape.top, slide_h, 720),
                "w_px": emu_to_px(shape.width, slide_w, 1280),
                "h_px": emu_to_px(shape.height, slide_h, 720),
                "x_emu": int(shape.left),
                "y_emu": int(shape.top),
                "w_emu": int(shape.width),
                "h_emu": int(shape.height),
            },
            "text": text,
            "style": style_summary(shape),
            "font": font_summary(shape),
            "source_xml_path": "ppt/slides/slide{}.xml".format(slide_index),
            "source_shape_id": str(shape.shape_id),
            "source_shape_name": getattr(shape, "name", ""),
        }
        if parent_group:
            obj["group_id"] = parent_group.get("id", "")
            obj["parent_group_name"] = parent_group.get("name", "")
            obj["group_depth"] = depth
        if obj["type"] == "table":
            obj["table"] = table_summary(shape)
        if obj["type"] == "image":
            try:
                obj["image_content_type"] = shape.image.content_type
                obj["image_ext"] = shape.image.ext
                obj["image_sha1"] = shape.image.sha1
            except Exception:
                pass
            rel = image_rels.get(str(shape.shape_id), {})
            obj["image_relationship_id"] = rel.get("relationship_id", "")
            obj["image_relationship_type"] = rel.get("relationship_type", "")
            obj["image_target"] = rel.get("target", "")
            obj["image_package_path"] = rel.get("package_path", "")
        elif image_rels.get(str(shape.shape_id)):
            rel = image_rels.get(str(shape.shape_id), {})
            obj["fill_image_relationship_id"] = rel.get("relationship_id", "")
            obj["fill_image_relationship_type"] = rel.get("relationship_type", "")
            obj["fill_image_target"] = rel.get("target", "")
            obj["fill_image_package_path"] = rel.get("package_path", "")
        if obj["type"] == "group" and hasattr(shape, "shapes"):
            try:
                obj["child_count"] = len(list(shape.shapes))
            except Exception:
                obj["child_count"] = 0
        return obj

    def append_shape_tree(slide_index: int, shape: Any, z: float, image_rels: Dict[str, Dict[str, str]], objects: List[Dict[str, Any]], parent_group: Optional[Dict[str, Any]] = None, depth: int = 0) -> None:
        obj = object_from_shape(slide_index, shape, z, image_rels, parent_group=parent_group, depth=depth)
        objects.append(obj)
        if obj["type"] == "group" and hasattr(shape, "shapes"):
            for child_index, child in enumerate(shape.shapes, start=1):
                append_shape_tree(slide_index, child, z + child_index / 1000.0, image_rels, objects, parent_group=obj, depth=depth + 1)

    for slide_index, slide in enumerate(prs.slides, start=1):
        image_rels = picture_relationships(pptx, slide_index)
        background = slide_background(pptx, slide_index)
        objects: List[Dict[str, Any]] = []
        for z, shape in enumerate(slide.shapes, start=1):
            append_shape_tree(slide_index, shape, float(z), image_rels, objects)
        object_count += len(objects)
        slides.append({
            "index": slide_index,
            "source_xml_path": "ppt/slides/slide{}.xml".format(slide_index),
            "background": background,
            "object_count": len(objects),
            "objects": objects,
        })
    return {
        "slide_count": len(slides),
        "object_count": object_count,
        "slide_size": {"w_emu": slide_w, "h_emu": slide_h, "w_px_base": 1280, "h_px_base": 720},
        "slides": slides,
    }


def extracted_text_markdown(object_inventory: Dict[str, Any]) -> str:
    lines = ["# Extracted Text", ""]
    for slide in object_inventory["slides"]:
        lines.append("## Slide {}".format(slide["index"]))
        has_text = False
        for obj in slide["objects"]:
            text = (obj.get("text") or "").strip()
            if text:
                has_text = True
                lines.append("- `{}`: {}".format(obj.get("name") or obj.get("id"), text.replace("\n", " / ")))
        if not has_text:
            lines.append("- _(no text)_")
        lines.append("")
    return "\n".join(lines)


def hex_color_from_node(node: ET.Element) -> Optional[str]:
    srgb = node.find(".//{}srgbClr".format(A_NS))
    if srgb is not None and srgb.attrib.get("val"):
        return srgb.attrib.get("val")
    sysclr = node.find(".//{}sysClr".format(A_NS))
    if sysclr is not None:
        return sysclr.attrib.get("lastClr") or sysclr.attrib.get("val")
    return None


def inventory_theme(unpacked: Path) -> Dict[str, Any]:
    themes: List[Dict[str, Any]] = []
    for theme_path in sorted((unpacked / "ppt" / "theme").glob("theme*.xml")):
        rel = str(theme_path.relative_to(unpacked)).replace("\\", "/")
        try:
            root = ET.fromstring(theme_path.read_bytes())
        except Exception as exc:
            themes.append({"path": rel, "parse_error": str(exc), "colors": {}, "fonts": {}})
            continue
        colors: Dict[str, str] = {}
        clr_scheme = root.find(".//{}clrScheme".format(A_NS))
        if clr_scheme is not None:
            for child in list(clr_scheme):
                color = hex_color_from_node(child)
                if color:
                    colors[child.tag.split("}")[-1]] = color
        fonts: Dict[str, str] = {}
        font_scheme = root.find(".//{}fontScheme".format(A_NS))
        if font_scheme is not None:
            for label, path in [("major_latin", ".//{}majorFont/{}latin"), ("minor_latin", ".//{}minorFont/{}latin")]:
                node = font_scheme.find(path.format(A_NS, A_NS))
                if node is not None and node.attrib.get("typeface"):
                    fonts[label] = node.attrib["typeface"]
        themes.append({"path": rel, "colors": colors, "fonts": fonts})
    return {"theme_count": len(themes), "themes": themes}


def rel_targets(rels_path: Path) -> List[Dict[str, str]]:
    if not rels_path.exists():
        return []
    try:
        root = ET.fromstring(rels_path.read_bytes())
    except Exception:
        return []
    out = []
    for rel in root:
        item = {k: v for k, v in rel.attrib.items()}
        out.append(item)
    return out


def inventory_masters_layouts(unpacked: Path) -> Dict[str, Any]:
    masters = []
    layouts = []
    master_dir = unpacked / "ppt" / "slideMasters"
    layout_dir = unpacked / "ppt" / "slideLayouts"
    for path in sorted(master_dir.glob("slideMaster*.xml")):
        rel = str(path.relative_to(unpacked)).replace("\\", "/")
        rels = path.parent / "_rels" / (path.name + ".rels")
        masters.append({"path": rel, "relationships": rel_targets(rels)})
    for path in sorted(layout_dir.glob("slideLayout*.xml")):
        rel = str(path.relative_to(unpacked)).replace("\\", "/")
        rels = path.parent / "_rels" / (path.name + ".rels")
        name = ""
        try:
            root = ET.fromstring(path.read_bytes())
            c_sld = root.find("{}cSld".format(P_NS))
            if c_sld is not None:
                name = c_sld.attrib.get("name", "")
        except Exception:
            pass
        layouts.append({"path": rel, "name": name, "relationships": rel_targets(rels)})
    return {"master_count": len(masters), "layout_count": len(layouts), "masters": masters, "layouts": layouts}


def inventory_assets(unpacked: Path) -> Dict[str, Any]:
    media = []
    media_dir = unpacked / "ppt" / "media"
    if media_dir.exists():
        for path in sorted(p for p in media_dir.rglob("*") if p.is_file()):
            rel = str(path.relative_to(unpacked)).replace("\\", "/")
            media.append({
                "path": rel,
                "size_bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "extension": path.suffix.lower().lstrip("."),
            })
    charts = []
    chart_dir = unpacked / "ppt" / "charts"
    if chart_dir.exists():
        for path in sorted(p for p in chart_dir.rglob("*") if p.is_file()):
            rel = str(path.relative_to(unpacked)).replace("\\", "/")
            charts.append({"path": rel, "size_bytes": path.stat().st_size, "sha256": sha256_file(path)})
    return {"media_count": len(media), "chart_count": len(charts), "media": media, "charts": charts}


def create_contact_sheet(rendered_dir: Path, thumb_width: int = 320, columns: int = 4) -> Path:
    """Create rendered/contact-sheet.png from slide-*.png files."""
    if Image is None:
        raise RuntimeError("Pillow is required to create a contact sheet")
    slides = sorted(p for p in rendered_dir.glob("slide-*.png") if p.is_file())
    if not slides:
        raise RuntimeError("no rendered slide PNGs found in {}".format(rendered_dir))
    columns = max(1, int(columns))
    margin = 18
    label_h = 26
    thumbs = []
    max_h = 0
    for idx, path in enumerate(slides, start=1):
        img = Image.open(str(path)).convert("RGB")
        ratio = thumb_width / float(img.width)
        thumb_h = max(1, int(round(img.height * ratio)))
        img = img.resize((thumb_width, thumb_h))
        thumbs.append((idx, path, img))
        max_h = max(max_h, thumb_h)
    rows = (len(thumbs) + columns - 1) // columns
    cell_w = thumb_width + margin
    cell_h = max_h + label_h + margin
    sheet = Image.new("RGB", (columns * cell_w + margin, rows * cell_h + margin), "white")
    draw = ImageDraw.Draw(sheet) if ImageDraw is not None else None
    for n, (idx, _path, img) in enumerate(thumbs):
        row = n // columns
        col = n % columns
        x = margin + col * cell_w
        y = margin + row * cell_h
        sheet.paste(img, (x, y))
        if draw is not None:
            draw.text((x, y + max_h + 4), "slide-{:02d}".format(idx), fill=(60, 60, 60))
    out = rendered_dir / "contact-sheet.png"
    sheet.save(str(out))
    return out


def render_if_requested(pptx: Path, outdir: Path, skip_render: bool, dpi: int) -> List[str]:
    if skip_render:
        return []
    rendered_dir = outdir / "rendered"
    try:
        sys.path.insert(0, str(ROOT / "scripts"))
        from render_pptx_png import render  # type: ignore
        written = render(pptx, rendered_dir, dpi=dpi)
        if written:
            create_contact_sheet(rendered_dir)
        return written
    except Exception as exc:
        (outdir / "render-error.txt").write_text(str(exc), encoding="utf-8")
        return []


def analyze(pptx: Path, outdir: Path, deck_id: Optional[str] = None, source_url: str = "", license_note: str = "", skip_render: bool = False, dpi: int = 96) -> Dict[str, Any]:
    pptx = pptx.resolve()
    if not pptx.exists():
        raise FileNotFoundError(str(pptx))
    clean_dir(outdir)
    original = outdir / "original.pptx"
    shutil.copyfile(str(pptx), str(original))
    unpacked = outdir / "unpacked"
    package_parts = unzip_pptx(original, unpacked)

    object_inventory = inventory_objects(original)
    theme_inventory = inventory_theme(unpacked)
    master_layout_inventory = inventory_masters_layouts(unpacked)
    asset_inventory = inventory_assets(unpacked)
    rendered = render_if_requested(original, outdir, skip_render=skip_render, dpi=dpi)

    provenance = {
        "deck_id": deck_id or pptx.stem,
        "original_filename": pptx.name,
        "original_path": str(pptx),
        "source_url": source_url,
        "license_note": license_note,
        "sha256": sha256_file(original),
        "created_at_utc": _dt.datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "slide_count": object_inventory["slide_count"],
        "object_count": object_inventory["object_count"],
        "theme_count": theme_inventory["theme_count"],
        "master_count": master_layout_inventory["master_count"],
        "layout_count": master_layout_inventory["layout_count"],
        "media_count": asset_inventory["media_count"],
        "package_part_count": len(package_parts),
        "rendered_slide_count": len(rendered),
        "render_skipped": bool(skip_render),
    }

    write_json(outdir / "provenance.json", provenance)
    write_json(outdir / "object-inventory.json", object_inventory)
    write_json(outdir / "theme-inventory.json", theme_inventory)
    write_json(outdir / "master-layout-inventory.json", master_layout_inventory)
    write_json(outdir / "asset-inventory.json", asset_inventory)
    (outdir / "extracted-text.md").write_text(extracted_text_markdown(object_inventory), encoding="utf-8")
    return provenance


def main() -> int:
    ap = argparse.ArgumentParser(description="Analyze PPTX specimen into clone-learning evidence pack")
    ap.add_argument("pptx", help="Input PPTX specimen")
    ap.add_argument("--out", required=True, help="Output specimen directory")
    ap.add_argument("--deck-id", default="", help="Stable deck/template id")
    ap.add_argument("--source-url", default="", help="Original source URL if known")
    ap.add_argument("--license-note", default="", help="License/attribution note")
    ap.add_argument("--skip-render", action="store_true", help="Skip LibreOffice/pdftoppm rendering")
    ap.add_argument("--dpi", type=int, default=96, help="Render DPI when rendering is enabled")
    args = ap.parse_args()
    try:
        summary = analyze(
            Path(args.pptx),
            Path(args.out),
            deck_id=args.deck_id or None,
            source_url=args.source_url,
            license_note=args.license_note,
            skip_render=args.skip_render,
            dpi=args.dpi,
        )
    except Exception as exc:
        print("FAIL analyze pptx specimen: {}".format(exc), file=sys.stderr)
        return 1
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
