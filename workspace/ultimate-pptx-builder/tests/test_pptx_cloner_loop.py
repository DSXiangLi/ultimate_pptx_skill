#!/usr/bin/env python3
"""Tests for the PPTX cloner acceptance loop."""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from pptx import Presentation

from tests.test_pptx_specimen_analyzer import make_sample_pptx

ROOT = Path(__file__).resolve().parents[1]
LOOP = ROOT / "scripts" / "run_pptx_cloner_loop.py"
sys.path.insert(0, str(ROOT / "scripts"))
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def slide_background_rgb(pptx: Path, slide_index: int = 1) -> str:
    with zipfile.ZipFile(str(pptx)) as zf:
        root = ET.fromstring(zf.read("ppt/slides/slide{}.xml".format(slide_index)))
    bg = root.find("{0}cSld/{0}bg/{0}bgPr/{1}solidFill/{1}srgbClr".format(P_NS, A_NS))
    return bg.attrib.get("val", "") if bg is not None else ""


class PptxClonerLoopTests(unittest.TestCase):
    def test_loop_passes_c1_c2_for_simple_deck_and_writes_report(self):
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            pptx = tmp / "sample.pptx"
            out = tmp / "specimen"
            report = out / "cloner-loop-report.json"
            make_sample_pptx(pptx)

            proc = subprocess.run(
                [
                    sys.executable,
                    str(LOOP),
                    str(pptx),
                    "--deck-id",
                    "sample-deck",
                    "--out",
                    str(out),
                    "--skip-render",
                    "--max-iterations",
                    "2",
                ],
                cwd=str(ROOT),
                text=True,
                capture_output=True,
            )
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            self.assertTrue(report.exists())
            data = json.loads(report.read_text(encoding="utf-8"))

            self.assertEqual(data["release_decision"], "pass")
            self.assertEqual(data["blocking_count"], 0)
            self.assertTrue(data["next_phase_allowed"])
            self.assertEqual(data["next_phase"], "C4 template archetype mining")
            self.assertTrue((out / "original.pptx").exists())
            self.assertTrue((out / "decompiled.raw.ir.json").exists())
            self.assertTrue((out / "rebuilt.pptx").exists())
            self.assertTrue((out / "rebuild-report.json").exists())
            self.assertEqual(slide_background_rgb(out / "rebuilt.pptx"), "123456")
            rebuild_report = json.loads((out / "rebuild-report.json").read_text(encoding="utf-8"))
            produced_by_type = {
                item["type"]: item["produced"]
                for item in rebuild_report["objects"]
                if item["type"] == "image"
            }
            self.assertIn("native-image", set(produced_by_type.values()))
            self.assertFalse(
                any(item["type"] == "image" for item in rebuild_report["unsupported_objects"]),
                "image objects with source media references should be reconstructed, not placeholdered",
            )
            grouped_child = next(item for item in rebuild_report["objects"] if item.get("parent_group_name") == "sample_group")
            group_container = next(item for item in rebuild_report["objects"] if item.get("type") == "group")
            self.assertEqual(group_container["produced"], "expanded-group-container")
            self.assertFalse(
                any(item["type"] == "group" for item in rebuild_report["unsupported_objects"]),
                "groups with emitted children should not remain opaque placeholders in C3",
            )
            self.assertIn(grouped_child["produced"], {"native-text", "native-shape"})
            empty_text = next(item for item in rebuild_report["objects"] if item.get("source_shape_name") == "sample_empty_text_container")
            self.assertEqual(empty_text["produced"], "skipped-empty-text-container")
            self.assertFalse(
                any(item["type"] == "text" for item in rebuild_report["unsupported_objects"]),
                "empty/decorative text containers should be classified and skipped, not placeholdered",
            )
            table = next(item for item in rebuild_report["objects"] if item.get("source_shape_name") == "sample_native_table")
            self.assertEqual(table["produced"], "native-table")
            self.assertGreaterEqual(table.get("styled_table_cells") or 0, 4)
            self.assertFalse(
                any(item["type"] == "table" for item in rebuild_report["unsupported_objects"]),
                "simple native tables should be reconstructed, not placeholdered",
            )
            rebuilt = Presentation(str(out / "rebuilt.pptx"))
            rebuilt_table = next(shape.table for shape in rebuilt.slides[0].shapes if getattr(shape, "has_table", False))
            rebuilt_header = rebuilt_table.cell(0, 0)
            rebuilt_run = rebuilt_header.text_frame.paragraphs[0].runs[0]
            self.assertEqual(str(rebuilt_header.fill.fore_color.rgb), "1A2B3C")
            self.assertEqual(rebuilt_run.font.name, "Aptos Display")
            self.assertEqual(round(float(rebuilt_run.font.size.pt)), 14)
            self.assertTrue(rebuilt_run.font.bold)
            self.assertEqual(str(rebuilt_run.font.color.rgb), "FFFFFF")
            gate_ids = {gate["id"] for gate in data["gates"]}
            self.assertIn("C1-EVIDENCE-PACK", gate_ids)
            self.assertIn("C2-TEXT-RECALL", gate_ids)
            self.assertIn("C2-CRITICAL-TEXT-EDITABILITY", gate_ids)
            self.assertIn("C3-REBUILD-PPTX", gate_ids)
            self.assertIn("C3-STRICT-PACKAGE", gate_ids)
            self.assertIn("C3-REBUILT-TEXT-RECALL", gate_ids)
            self.assertIn("c3_artifacts", data)

    def test_loop_materializes_picture_fill_shapes_as_native_images(self):
        pptx = ROOT / "research" / "pptx-template-library" / "files" / "it-software-sales-proposal-slides.pptx"
        if not pptx.exists():
            self.skipTest("template library fixture is unavailable")
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "specimen"
            proc = subprocess.run(
                [
                    sys.executable,
                    str(LOOP),
                    str(pptx),
                    "--deck-id",
                    "it-software-sales-proposal-slides",
                    "--out",
                    str(out),
                    "--skip-render",
                    "--max-iterations",
                    "2",
                ],
                cwd=str(ROOT),
                text=True,
                capture_output=True,
            )
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            rebuild_report = json.loads((out / "rebuild-report.json").read_text(encoding="utf-8"))
            shape315 = next(item for item in rebuild_report["objects"] if item.get("source_shape_id") == "315")
            self.assertEqual(shape315["produced"], "native-picture-fill-shape-image")

    def test_evaluator_fails_with_actionable_queue_for_missing_ir(self):
        from run_pptx_cloner_loop import evaluate_c1_c2  # type: ignore
        from analyze_pptx_specimen import analyze  # type: ignore

        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            pptx = tmp / "sample.pptx"
            out = tmp / "specimen"
            make_sample_pptx(pptx)
            analyze(pptx, out, deck_id="sample-deck", skip_render=True)

            report = evaluate_c1_c2(out, require_render=False)

            self.assertEqual(report["release_decision"], "fail")
            self.assertGreater(report["blocking_count"], 0)
            codes = {item["code"] for item in report["optimization_queue"]}
            self.assertIn("decompiled_ir_missing", codes)
            owners = {item["owner"] for item in report["optimization_queue"]}
            self.assertIn("C2 decompiler", owners)


if __name__ == "__main__":
    unittest.main()
