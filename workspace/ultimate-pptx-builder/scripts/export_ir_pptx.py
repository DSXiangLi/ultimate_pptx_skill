#!/usr/bin/env python3
"""Export Slide IR to an editable Office-compatible PPTX.

The previous exporter hand-wrote a minimal OOXML package. LibreOffice and
python-pptx could open it, but Windows PowerPoint/WPS may still show repair
warnings for subtle package/schema details. This exporter uses python-pptx as
the package writer and only controls the editable drawing objects.
"""
from __future__ import annotations

from pathlib import Path
import argparse
import html
import json
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000
BASE_W_PX = 1280
BASE_H_PX = 720
EMU_PER_PX_X = SLIDE_W_EMU / BASE_W_PX
EMU_PER_PX_Y = SLIDE_H_EMU / BASE_H_PX
CRITICAL_ROLES = {"title", "body", "risk", "source", "footnote", "metric", "metric-label", "metric-note"}


def esc(s):
    return html.escape(str(s), quote=True)


def emu(v, axis="x"):
    return Emu(int(round(float(v) * (EMU_PER_PX_X if axis == "x" else EMU_PER_PX_Y))))


def rgb(value, default="111827"):
    v = str(value or default).strip().lstrip("#").upper()
    if len(v) == 3:
        v = "".join(ch * 2 for ch in v)
    if len(v) != 6 or any(ch not in "0123456789ABCDEF" for ch in v):
        v = default
    return RGBColor(int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))


def set_transparency(fmt, transparency):
    try:
        fmt.transparency = max(0.0, min(1.0, float(transparency)))
    except Exception:
        pass


def box_args(box):
    return emu(box["x"], "x"), emu(box["y"], "y"), emu(box["w"], "x"), emu(box["h"], "y")


def mso_shape(name):
    return {
        "rect": MSO_SHAPE.RECTANGLE,
        "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
        "ellipse": MSO_SHAPE.OVAL,
    }.get(name or "rect", MSO_SHAPE.RECTANGLE)


def safe_name(value, fallback):
    text = str(value or fallback).strip()
    return text if text else fallback


def add_shape(slide, obj, shape_id_hint=None, placeholder=False):
    box = obj["box"]
    shp = slide.shapes.add_shape(mso_shape(obj.get("shape")), *box_args(box))
    shp.name = safe_name(obj.get("id") + ("__raster_placeholder" if placeholder else ""), f"shape_{shape_id_hint or len(slide.shapes)}")
    fill = shp.fill
    fill.solid()
    fill.fore_color.rgb = rgb(obj.get("fill"), "E5E7EB")
    set_transparency(fill, 1.0 - float(obj.get("opacity", 1.0)))
    line = shp.line
    if float(obj.get("stroke_opacity", 1.0)) <= 0:
        line.fill.background()
    else:
        line.color.rgb = rgb(obj.get("stroke"), "D1D5DB")
        try:
            line.width = Pt(0.75)
        except Exception:
            pass
        set_transparency(line.fill, 1.0 - float(obj.get("stroke_opacity", 1.0)))
    return shp


def line_spacing_for(role, size_pt):
    if role in {"risk", "source", "footnote"}:
        return 1.18
    if role in {"metric", "metric-label", "metric-note"}:
        return 1.12
    if size_pt <= 9:
        return 1.26
    return 1.30


def add_text(slide, obj, shape_id_hint=None):
    box = obj["box"]
    shp = slide.shapes.add_textbox(*box_args(box))
    shp.name = safe_name(obj.get("id"), f"text_{shape_id_hint or len(slide.shapes)}")
    st = obj.get("style", {})
    role = obj.get("role", "")
    size_pt = float(st.get("size", 16))
    tf = shp.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    try:
        p.line_spacing = line_spacing_for(role, size_pt)
    except Exception:
        pass
    run = p.add_run()
    run.text = str(obj.get("text", ""))
    font = run.font
    font.name = st.get("font", "Microsoft YaHei")
    font.size = Pt(size_pt)
    font.bold = bool(st.get("bold"))
    font.color.rgb = rgb(st.get("color", "111827"), "111827")
    return shp


def make_text_obj(oid, text, x, y, w, h, size=10, color_value="B6C7D8", bold=False, font="Microsoft YaHei"):
    return {
        "id": oid,
        "type": "text",
        "role": "chart-label",
        "box": {"x": x, "y": y, "w": w, "h": h},
        "text": text,
        "style": {"font": font, "size": size, "color": color_value, "bold": bold},
        "editability": {"priority": 4},
    }


def make_shape_obj(oid, x, y, w, h, fill, stroke=None, opacity=1.0, stroke_opacity=0.0, prst="rect"):
    return {
        "id": oid,
        "type": "shape",
        "role": "chart-mark",
        "box": {"x": x, "y": y, "w": w, "h": h},
        "fill": fill,
        "stroke": stroke or fill,
        "opacity": opacity,
        "stroke_opacity": stroke_opacity,
        "shape": prst,
    }


def chart_tone_color(tone, style=None, index=0):
    style = style or {}
    palette = style.get("series_palette") or []
    if palette:
        return palette[index % len(palette)]
    return {"positive": "31D0AA", "warning": "FBBF24", "accent": "7C3AED"}.get(tone, "5ED7FF")


def add_chart_vector(slide, obj):
    box = obj["box"]
    x, y, w, h = box["x"], box["y"], box["w"], box["h"]
    st = obj.get("style", {})
    categories = obj.get("categories", [])
    series = obj.get("series", [])
    created = 0

    def s(oid, sx, sy, sw, sh, fill, stroke=None, opacity=1.0, stroke_opacity=0.0, prst="rect"):
        nonlocal created
        add_shape(slide, make_shape_obj(oid, sx, sy, sw, sh, fill, stroke, opacity, stroke_opacity, prst), created)
        created += 1

    def t(oid, text, tx, ty, tw, th, size=10, color_value="B6C7D8", bold=False, font=None):
        nonlocal created
        add_text(slide, make_text_obj(oid, text, tx, ty, tw, th, size, color_value, bold, font or st.get("font", "Microsoft YaHei")), created)
        created += 1

    s(obj["id"] + "_panel", x, y, w, h, st.get("panel_fill", "13243A"), st.get("panel_stroke", "2A6F91"), st.get("opacity", 0.48), st.get("stroke_opacity", 0.75), "roundRect")
    t(obj["id"] + "_title", obj.get("title", ""), x + 18, y + 8, max(1, w - 36), 20, 12, st.get("title_color", "EAF7FF"), True, st.get("title_font"))
    for idx, ser in enumerate(series[:3]):
        gx = x + 18 + idx * 72
        c = chart_tone_color(ser.get("tone"), st, idx)
        s(obj["id"] + "_legend_swatch_%d" % (idx + 1), gx, y + 32, 10, 10, c, c, 0.92, 0.0)
        t(obj["id"] + "_legend_label_%d" % (idx + 1), ser.get("name", ""), gx + 15, y + 29, 54, 14, 9, st.get("label_color", "B6C7D8"), False)
    plot_x, plot_y, plot_w, plot_h = x + 44, y + 66, max(1, w - 70), max(1, h - 116)
    max_total = 100
    if categories:
        for i in range(len(categories)):
            total = sum(float(s0.get("values", [0] * len(categories))[i]) for s0 in series if i < len(s0.get("values", [])))
            max_total = max(max_total, total)
    for val in [0, 50, 100]:
        yy = plot_y + plot_h - plot_h * val / max_total
        s(obj["id"] + "_grid_%d" % val, plot_x, yy, plot_w, 1, st.get("grid_color", "2A6F91"), st.get("grid_color", "2A6F91"), st.get("grid_opacity", 0.22), 0.0)
        t(obj["id"] + "_axis_%d" % val, "%d%%" % val, x + 8, yy - 7, 32, 12, 8, st.get("label_color", "B6C7D8"), False)
    if categories:
        gap = 10
        bw = max(10, (plot_w - gap * (len(categories) + 1)) / len(categories))
        bar_prst = st.get("bar_shape", "rect") if st.get("bar_shape") in {"rect", "roundRect"} else "rect"
        for i, cat in enumerate(categories):
            bx = plot_x + gap + i * (bw + gap)
            cursor = plot_y + plot_h
            for sidx, ser in enumerate(series):
                vals = ser.get("values", [])
                val = float(vals[i]) if i < len(vals) else 0
                bh = plot_h * val / max_total
                cursor -= bh
                c = chart_tone_color(ser.get("tone"), st, sidx)
                s(obj["id"] + "_bar_%d_%d" % (i + 1, sidx + 1), bx, cursor, bw, max(1, bh), c, c, 0.88, 0.0, bar_prst)
                if i == len(categories) - 1 and bh >= 13:
                    t(obj["id"] + "_label_%d_%d" % (i + 1, sidx + 1), str(int(round(val))), bx, cursor + bh / 2 - 6, bw, 12, 8, st.get("value_label_color", "EAF7FF"), False)
            t(obj["id"] + "_cat_%d" % (i + 1), str(cat), bx - 4, y + h - 38, bw + 8, 14, 9, st.get("label_color", "B6C7D8"), False)
    source = str(obj.get("source", ""))[:42]
    if source:
        t(obj["id"] + "_source", source, x + 18, y + h - 18, w - 36, 12, 7, st.get("label_color", "B6C7D8"), False)


def add_table_vector(slide, obj):
    box = obj["box"]
    x, y, w, h = box["x"], box["y"], box["w"], box["h"]
    st = obj.get("style", {})
    cols = obj.get("columns", [])
    rows = obj.get("rows", [])

    def s(oid, sx, sy, sw, sh, fill, stroke=None, opacity=1.0, stroke_opacity=0.0):
        add_shape(slide, make_shape_obj(oid, sx, sy, sw, sh, fill, stroke, opacity, stroke_opacity))

    def t(oid, text, tx, ty, tw, th, size=9, color_value="EAF7FF", bold=False):
        add_text(slide, make_text_obj(oid, text, tx, ty, tw, th, size, color_value, bold))

    s(obj["id"] + "_panel", x, y, w, h, st.get("row_fill", "13243A"), st.get("stroke", "2A6F91"), 0.52, st.get("stroke_opacity", 0.75))
    t(obj["id"] + "_title", obj.get("title", ""), x + 14, y + 10, w - 28, 20, 13, st.get("title_color", "EAF7FF"), True)
    if cols:
        tx, ty = x + 14, y + 38
        tw, th = w - 28, h - 64
        col_w = tw / len(cols)
        row_h = th / max(1, len(rows) + 1)
        for c, col in enumerate(cols):
            cx = tx + c * col_w
            s(obj["id"] + "_head_%d" % c, cx, ty, col_w, row_h, st.get("header_fill", "19324D"), st.get("stroke", "2A6F91"), 0.82, 0.5)
            t(obj["id"] + "_head_text_%d" % c, str(col)[:18], cx + 6, ty + 6, col_w - 12, row_h - 10, 10, st.get("label_color", "B6C7D8"), True)
        for r, row in enumerate(rows):
            ry = ty + (r + 1) * row_h
            fill = st.get("row_fill", "13243A") if r % 2 == 0 else st.get("row_alt_fill", "102135")
            for c, cell in enumerate(row[:len(cols)]):
                cx = tx + c * col_w
                s(obj["id"] + "_cell_%d_%d" % (r, c), cx, ry, col_w, row_h, fill, st.get("stroke", "2A6F91"), 0.58, 0.35)
                t(obj["id"] + "_cell_text_%d_%d" % (r, c), str(cell), cx + 6, ry + 5, col_w - 12, row_h - 10, 9, "EAF7FF", False)
    source = str(obj.get("source", ""))[:72]
    if source:
        t(obj["id"] + "_source", source, x + 14, y + h - 16, w - 28, 12, 7, st.get("label_color", "B6C7D8"), False)


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            return layout
    return prs.slide_layouts[6]


def build_pptx(ir, pptx_path, report_path):
    deck = ir["deck"]
    slides = deck.get("slides", [])
    if not slides:
        raise ValueError("deck has no slides")
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    layout = blank_layout(prs)
    audit = []
    for slide_ir in slides:
        slide = prs.slides.add_slide(layout)
        objects = sorted(slide_ir.get("objects", []), key=lambda o: o.get("z", 0))
        for obj in objects:
            priority = obj.get("editability", {}).get("priority", 1)
            role = obj.get("role", "")
            expected = "native" if priority >= 4 or role in CRITICAL_ROLES else obj.get("render_policy")
            actual = "unsupported"
            passed = False
            note = ""
            typ = obj.get("type")
            if typ == "text":
                add_text(slide, obj)
                actual = "native-text"
                passed = True
            elif typ == "shape":
                add_shape(slide, obj)
                actual = "native-shape"
                passed = True
            elif typ == "chart":
                add_chart_vector(slide, obj)
                actual = "editable-vector-chart"
                passed = True
                note = "Chart exported as editable vector group of native shapes/text."
            elif typ == "table":
                add_table_vector(slide, obj)
                actual = "editable-vector-table"
                passed = True
                note = "Table exported as editable vector group of native shapes/text."
            elif typ == "rasterIsland":
                placeholder_obj = dict(obj)
                placeholder_obj["fill"] = "F3F4F6"
                placeholder_obj["stroke"] = "E5E7EB"
                placeholder_obj["shape"] = "rect"
                add_shape(slide, placeholder_obj, placeholder=True)
                actual = "native-placeholder-for-rasterIsland"
                passed = priority <= 2 and role not in CRITICAL_ROLES and "source" in obj
                note = "Raster island exported as replaceable placeholder."
            else:
                note = "unsupported object type %s" % typ
            audit.append({
                "slide_id": slide_ir.get("id"),
                "id": obj.get("id"),
                "role": role,
                "type": typ,
                "priority": priority,
                "expected": expected,
                "actual": actual,
                "pass": passed,
                "note": note,
            })
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))
    blocking = [o for o in audit if (o["priority"] >= 4 or o["role"] in CRITICAL_ROLES) and not o["pass"]]
    report = {
        "deck_id": deck.get("id"),
        "pptx_path": str(pptx_path),
        "exporter": "python-pptx",
        "objects_audited": audit,
        "blocking_issues": blocking,
        "release_decision": "pass" if not blocking else "fail",
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help="Slide IR JSON")
    ap.add_argument("output", help="PPTX output path")
    ap.add_argument("--report", default=None, help="export report JSON path")
    args = ap.parse_args()
    ir = json.loads(Path(args.input).read_text(encoding="utf-8"))
    report = Path(args.report) if args.report else Path(args.output).with_suffix(".export-report.json")
    build_pptx(ir, Path(args.output), report)
    print("wrote %s" % args.output)
    print("wrote %s" % report)


if __name__ == "__main__":
    main()
